"""Build the local progress report from supplied evidence; no network or live tests."""
import sys
from pathlib import Path
import json
import hashlib

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'docs/.work-speech/lib'))
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE = ROOT / 'Dokumentasi/Progres 2'
ASSETS = ROOT / 'docs/progres2-assets'
ASSETS.mkdir(exist_ok=True)
DEST = SOURCE / 'progres 2 Try and error sistem.docx'
INK = '183B43'
TEAL = '286C68'
MUTED = '56666D'
LIGHT = 'E9F3F0'

# Irreversible masks are applied only to embedded copies, never to source images.
MASKS = {
 'dokumentasi penuh semua bot.png': [(40,790,260,822),(155,828,470,852),(40,914,260,947),(155,950,470,976),(40,1030,260,1063),(155,1066,470,1094),(40,1145,260,1179),(155,1182,505,1208)],
 'menu akun pengguna.png': [(24,364,177,1105)],
 'menu database dokumentasi llm.png': [(25,457,204,632)],
 'menu database log bot BPJS.png': [(25,457,215,632)],
 'menu database segmen transkrip.png': [(25,457,215,949)],
 'menu database log pengguna.png': [(25,457,239,1582)],
 'menu log admin.png': [(397,537,632,665)],
 'percobaan koneksi VPN data asli.jpeg': [(190,760,465,794),(190,943,465,978),(55,1120,248,1230)],
 'try&error pada vpn.jpeg': [(140,471,361,501),(145,608,366,639)],
 'try&error pada termina.jpeg': [(96,667,377,713),(96,803,350,850),(264,1207,406,1235)],
 'percobaan koneksi terminal data asli.jpeg': [(133,104,555,153),(0,918,708,980),(0,984,195,1017)],
}
font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 14)
manifest = []
for path in sorted(SOURCE.iterdir()):
    if path.suffix.lower() not in ('.png','.jpeg','.jpg'):
        continue
    im = Image.open(path).convert('RGB')
    draw = ImageDraw.Draw(im)
    for x1,y1,x2,y2 in MASKS.get(path.name, []):
        draw.rectangle((x1,y1,x2,y2), fill='#d6dddd')
        if x2-x1 > 110 and y2-y1 > 23:
            draw.text((x1+5,y1+5), 'Disamarkan', fill='#45545a', font=font)
    target = ASSETS / (path.stem + '.png')
    im.save(target)
    manifest.append({'source': path.name, 'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'report_asset': target.name, 'masked_regions': MASKS.get(path.name, [])})
if manifest:
    (ASSETS/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
else:
    # A checkout can contain only the sanitized assets; originals are private.
    manifest=json.loads((ASSETS/'manifest.json').read_text(encoding='utf-8'))
(ASSETS/'README.md').write_text('''# Gambar laporan Progres 2

21 salinan gambar dari `Dokumentasi/Progres 2`, dibuat untuk laporan Word 19 Oktober 2026. Sumber asli tidak diubah. Bidang identitas/akun/ID/alamat server ditutup permanen pada salinan, tanpa mengganti status keberhasilan/kegagalan pada layar. `manifest.json` mencatat nama sumber, hash, dan area penyamaran. Bukan bukti bahwa data asal seluruhnya fiktif.

Gambar adalah dokumentasi tampilan; keluaran suara, kecepatan, akurasi STT/LLM, dan pengiriman lintas perangkat memerlukan uji tersendiri. Pembuat laporan: `docs/.work-progres2/build_report.py`.
''',encoding='utf-8')

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.85), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2.1), Cm(2.1)
sec.header_distance, sec.footer_distance = Cm(.8), Cm(.8)
sec.different_first_page_header_footer = True
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(7)
for name, size, color in [('Title',32,INK),('Heading 1',21,INK),('Heading 2',13,TEAL),('Heading 3',11,TEAL)]:
    s=doc.styles[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
    s.paragraph_format.space_before=Pt(9); s.paragraph_format.space_after=Pt(8)
caption=doc.styles['Caption']; caption.font.name='Calibri'; caption.font.size=Pt(9); caption.font.color.rgb=RGBColor.from_string(MUTED); caption.font.italic=False
header=sec.header.paragraphs[0]
header.text='AI HUB  /  LAPORAN PERKEMBANGAN SISTEM'
header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor.from_string(TEAL)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('PROGRES 2  •  19 OKTOBER 2026     |     ').font.size=Pt(8)
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); footer._p.append(f)
doc.core_properties.title='Progres 2 — Try and Error Sistem AI Hub'
doc.core_properties.subject='Laporan perkembangan prototipe, penambahan komunikasi, dan hasil percobaan'
doc.core_properties.author='Khoirul Yardan'
doc.core_properties.keywords='AI Hub, Progres 2, perawat-pasien, Papan Komunikasi, Anamnesis Terpandu'
doc.core_properties.comments='Disusun dari bukti lokal; pengujian aplikasi tidak dijalankan ulang saat penyusunan.'

pages=[]
figures=[]
def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text):
    doc.add_heading(text,2)

def page(title, tag=None):
    if pages: doc.add_page_break()
    pages.append(title)
    if tag:
        t=p(tag.upper()); t.runs[0].font.size=Pt(9); t.runs[0].font.bold=True; t.runs[0].font.color.rgb=RGBColor.from_string(TEAL)
    doc.add_heading(title,1)

def box(title, text):
    t=doc.add_table(rows=1,cols=1)
    cell=t.cell(0,0)
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),LIGHT); tcPr.append(sh)
    r=cell.paragraphs[0].add_run(title); r.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL)
    cell.add_paragraph(text)
    p('') .paragraph_format.space_after=Pt(0)

def table(headers, rows, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.autofit=False
    for i, text in enumerate(headers):
        cell=t.rows[0].cells[i]; cell.text=text
        sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),TEAL); cell._tc.get_or_add_tcPr().append(sh)
        for run in cell.paragraphs[0].runs: run.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    rep=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
    for idx, row in enumerate(rows):
        cells=t.add_row().cells
        for i,text in enumerate(row):
            cells[i].text=str(text)
            if idx%2==0:
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'F2F6F5'); cells[i]._tc.get_or_add_tcPr().append(sh)
    for row in t.rows:
        no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for i,cell in enumerate(row.cells):
            if widths: cell.width=Cm(widths[i])
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(5); para.paragraph_format.space_before=Pt(4); para.paragraph_format.line_spacing=1.0
                for run in para.runs: run.font.size=Pt(9.5)
    p('').paragraph_format.space_after=Pt(0)
    return t

def add_image(paragraph, name, max_width, max_height):
    path=ASSETS/(Path(name).stem+'.png')
    w,ht=Image.open(path).size
    width=min(max_width,max_height*w/ht)
    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after=Pt(4)
    paragraph.paragraph_format.keep_with_next=True
    shape=paragraph.add_run().add_picture(str(path),width=Inches(width))
    shape._inline.docPr.set('descr',name+'; salinan laporan, identitas disamarkan bila diperlukan')

def figure(name, explanation, max_height=5.3):
    add_image(p(''),name,5.8,max_height)
    num=len(figures)+1; figures.append((num,name,explanation))
    cap=p(f'Gambar {num}. {explanation}', 'Caption'); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER

def pair(items,max_height=4.8):
    t=doc.add_table(rows=1,cols=2); t.autofit=False
    for cell,(name,explanation) in zip(t.rows[0].cells,items):
        cell.width=Cm(8.2)
        add_image(cell.paragraphs[0],name,3.0,max_height)
        num=len(figures)+1; figures.append((num,name,explanation))
        cap=cell.add_paragraph(f'Gambar {num}. {explanation}','Caption'); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    no=OxmlElement('w:cantSplit'); t.rows[0]._tr.get_or_add_trPr().append(no)
    p('').paragraph_format.space_after=Pt(0)

# Page 1: cover
pages.append('Sampul')
p('LAPORAN PERKEMBANGAN PROYEK AKHIR').runs[0].font.color.rgb=RGBColor.from_string(TEAL)
p('\n')
doc.add_heading('PROGRES 2',0)
doc.add_heading('Try and Error\nSistem AI Hub',0)
p('Pengembangan fitur komunikasi, dokumentasi,\ndan pengujian infrastruktur', 'Subtitle')
p('\n')
box('Fokus laporan','Papan Komunikasi • Anamnesis Terpandu • Dokumentasi & Riwayat • VPN dan Terminal • Panel Admin')
p('Disusun oleh: Khoirul Yardan')
p('Proyek Akhir D3 PENS Lamongan')
p('Tanggal penyusunan: 19 Oktober 2026')
h('Judul Proyek Akhir')
p('Pengembangan Aplikasi Mobile untuk Otomatisasi Dokumentasi Percakapan Perawat-Pasien sebagai Pendukung Verifikasi Diagnosis Pasien Menggunakan Speech-to-Text dan Large Language Model.')
p('Laporan ini mencatat perkembangan prototipe dan proses percobaan berdasarkan dokumentasi yang tersedia. Penambahan dua modul komunikasi dilakukan secara mandiri sambil menunggu tanggapan dosen atas pengajuan sebelumnya.')

page('Panduan membaca dan daftar isi','Orientasi laporan')
p('Laporan disusun dari dokumen awal, catatan percobaan bertanggal, kode lokal, dan 21 tangkapan layar dalam folder Progres 2. Bentuk penjelasan mengikuti dokumen “Dokumentasi macam speech perawat dan pasien”, tetapi gambar dan pembahasannya diarahkan pada perkembangan sistem yang benar-benar terdokumentasi.')
toc=table(['Bagian','Pokok pembahasan'],[
 ('1. Ringkasan perkembangan','Hasil dan batas bukti'),('2. Kronologi pengembangan','Perubahan dari awal sampai Progres 2'),('3. Speech dan posisi tiap modul','Pemetaan 9 jenis dan batas akademik'),('4. Papan Komunikasi','Pilihan, ruangan/tautan, dan catatan'),('5. Anamnesis Terpandu','Pertanyaan sistem dan jawaban pasien'),('6. Dokumentasi & Riwayat','Pencarian dan tiga sumber data'),('7–8. Try and error','OpenVPN dan Terminal SSH'),('9. Panel Admin','Login, dashboard, dan akun pengguna'),('9.2–9.3. Metadata database','Sesi, draf, segmen, dan aktivitas'),('9.4–9.6. Pengelolaan sistem','Maintenance, pengumuman, audit/status'),('10–11. Evaluasi dan tindak lanjut','Rencana pengujian serta penutup'),('Lampiran','Sumber, daftar gambar, penataan berkas')
],[8,8.3])
h('Arti status dalam laporan')
p('Terlihat pada gambar: keadaan antarmuka saat diambil. Tersedia di kode: implementasi ditemukan, belum tentu lulus uji lengkap. Tercatat lulus uji: hasil pada lingkungan/tanggal tertentu menurut catatan sumber. Belum dibuktikan: tidak ada bukti yang cukup dalam bahan laporan.')
p('Identitas pada riwayat, data akun, ID terkait, dan alamat server disamarkan di salinan gambar. Gambar asli tetap dipertahankan. Contoh dialog yang ditulis dalam laporan adalah fiktif; laporan tidak mengasumsikan seluruh identitas pada sumber asli merupakan data fiktif.')

page('1. Ringkasan perkembangan','Gambaran hasil')
p('Progres 2 melanjutkan aplikasi AI Hub yang sebelumnya telah memiliki Bot BPJS untuk dokumentasi percakapan, Chat, VPN, dan Terminal. Perkembangan kali ini menambahkan jalur komunikasi melalui Papan Komunikasi dan Anamnesis Terpandu, menyatukan akses dokumentasi, serta menyediakan panel admin untuk pengelolaan operasional.')
p('Pengajuan pemilihan speech sebelumnya belum memperoleh jawaban dosen menurut keterangan pengguna. Oleh karena itu, kedua tambahan dilaporkan sebagai pengembangan mandiri yang masih perlu dibahas. Ketersediaan fitur tidak ditulis sebagai persetujuan pembimbing atau bukti keberhasilan klinis.')
table(['Area','Hasil yang dapat dilaporkan','Batas bukti'],[
 ('Papan Komunikasi','Pilihan kebutuhan, catatan, ruangan/tautan terlihat.','Suara dan pengiriman dua perangkat perlu uji langsung.'),
 ('Anamnesis Terpandu','Alur pertanyaan TTS, jawaban STT, koreksi dan draf tersedia di kode.','Riwayat pada gambar menunjukkan 0 sesi Anamnesis.'),
 ('Dokumentasi','Filter tiga sumber tersedia; gambar berisi 7 entri.','Jumlah entri bukan ukuran ketepatan isi.'),
 ('OpenVPN','Uji 18 Oktober berhasil setelah penyesuaian negosiasi cipher.','Baru lingkungan Wi-Fi yang diuji.'),
 ('Terminal SSH','Server asli terbuka sampai prompt setelah VPN aktif.','Perintah di server asli belum diuji.'),
 ('Panel Admin','Dashboard, akun, metadata, maintenance, pengumuman, audit/status.','Tidak menampilkan isi klinis atau memverifikasi diagnosis.')
],[3,7.3,6])
box('Hasil utama','Prototipe bertambah lengkap untuk dicoba dan didiskusikan. Evaluasi akurasi speech, kelengkapan draf, konsistensi isi, serta waktu proses tetap membutuhkan eksperimen terukur.')
p('Penyusunan laporan tidak menjalankan ulang aplikasi, tidak memanggil provider berbayar, dan tidak mengubah database. Hasil pengujian sebelumnya dipakai sesuai tanggal dan batas lingkungan yang tertulis pada sumber.')

page('2. Kronologi pengembangan','Dari dokumentasi awal sampai akhir')
table(['Tahap','Perkembangan','Dasar / batas'],[
 ('September 2026','Kajian kelayakan hub AI, API dan arsitektur hybrid.','Dokumen kemungkinan pembangunan; masih mencakup gagasan lama.'),
 ('2–5 Oktober','Perancangan dan implementasi UI; penyederhanaan Bot BPJS.','Desain, implementasi, dan snapshot sistem lama.'),
 ('Progres 1','Gambar STT, ringkasan, perbandingan, tukar peran, chat, VPN/Terminal.','16 gambar tahap sebelumnya; tanggal tiap gambar tidak ditebak.'),
 ('14 Oktober','Uji web/Docker dan VPN/SSH tiruan; percobaan chat.','Jalur sukses chat disimulasikan karena percobaan provider nyata terkendala kuota.'),
 ('15–16 Oktober','Kajian/pengajuan speech; Papan dan Anamnesis dipilih sebagai tambahan.','Skor pemilihan adalah penilaian usulan, bukan hasil eksperimen.'),
 ('17 Oktober','Pengujian web, RLS tiruan, Flutter; audit BPJS dan relabel.','Angka hasil berlaku pada lingkungan sumber.'),
 ('18 Oktober','Perbaikan cipher OpenVPN; SSH asli sampai prompt.','Catatan percobaan perangkat; Wi-Fi sama dengan laptop.'),
 ('18 Oktober','Penyederhanaan admin, deploy Fly, uji desktop dan Chrome HP.','HANDOFF; bukan uji semua alur native speech.'),
 ('Progres 2','Bukti papan, riwayat, admin, serta gagal/berhasil VPN dan Terminal.','21 gambar, dirangkum 19 Oktober 2026.')
],[3,7.4,5.9])
h('Penyesuaian terhadap dokumentasi lama')
p('Pernyataan lama bahwa diarization belum tersedia, OpenVPN belum memiliki tunnel nyata, dan admin masih mengelola katalog bot/provider sudah tidak menggambarkan keadaan terakhir. Versi lama tetap disimpan sebagai arsip agar proses perubahan dapat ditelusuri.')
p('Verifikasi dokter tetap berlangsung di luar aplikasi. Penghapusan Friend System, wake word, dan balasan suara Bot BPJS tidak dibatalkan. TTS pada Papan dan Anamnesis memiliki fungsi berbeda dari balasan suara Bot BPJS.')

page('3. Speech dan posisi tiap modul','Dasar pemilihan dan keselarasan PA')
p('Dokumen referensi memuat sembilan bentuk komunikasi. Daftar itu merupakan peta konsep untuk diskusi, sehingga tidak semua jenis dapat disebut telah diimplementasikan. Berikut penempatannya pada keadaan sistem yang ditinjau.')
table(['Jenis komunikasi','Penerapan / status'],[
 ('1. Wawancara terapeutik','Percakapan perawat–pasien yang didokumentasikan melalui Bot BPJS.'),
 ('2. STT untuk dokumentasi','Bot BPJS; jawaban suara pada Anamnesis juga menggunakan STT.'),
 ('3. Caption langsung untuk pasien','Belum menjadi fitur yang dibuktikan pada Progres 2.'),
 ('4. Papan gambar/simbol','Papan Komunikasi tersedia dan terlihat pada gambar.'),
 ('5. Teks menjadi suara','TTS pada pilihan papan dan pertanyaan Anamnesis; ketik bebas belum dibuktikan sebagai fitur.'),
 ('6. Ya / Tidak / Tidak tahu','Masih backlog, bukan fitur selesai pada bukti ini.'),
 ('7. Lokasi dan intensitas nyeri','Tersedia dalam alur nyeri pada kode Papan; uji lengkap pemilihan perlu dicatat.'),
 ('8. Teach-back','Masih gagasan pengembangan, belum fitur selesai.'),
 ('9. Akses tatapan mata','Tidak diimplementasikan; di luar cakupan prototipe saat ini.')
],[5.1,11.2])
h('Bot BPJS tetap menjadi inti')
p('Audio percakapan perawat–pasien diubah menjadi transkrip dengan STT dan diarization. LLM membantu ekstraksi serta penyusunan draf; perawat memeriksa dan menyunting sebelum membagikannya. Peran pembicara berasal dari konteks sesi, bukan pencocokan biometrik suara.')
box('Dua tambahan mandiri','Papan Komunikasi memberi jalur pilihan ikon/TTS. Anamnesis Terpandu memakai pertanyaan sistem/TTS dan jawaban pasien/STT. Keduanya mendukung pengumpulan informasi; tidak mengganti objek utama penelitian, tidak merekam dokter, dan tidak menetapkan diagnosis.')

page('4. Papan Komunikasi','Pilihan kebutuhan pasien')
p('Papan menyediakan pilihan sentuh untuk menyampaikan kebutuhan tanpa harus menyusun ucapan panjang. Pilihan yang terlihat meliputi Nyeri, Sesak napas, Haus, Lapar, Mual, Ke toilet, Dingin, Panas, Ubah posisi, Minta obat, dan Panggil Suster.')
figure('papan komunikasi.png','Papan Komunikasi dengan 11 pilihan kebutuhan dan tab Papan, Ruangan, Catatan.',4.9)
h('Cara kerja dan pengamatan')
p('Pilihan menjadi kalimat TTS dan event catatan. Pada nyeri, kode menyediakan lokasi tubuh dan skala 0–10 sebagai laporan pasien. Perawat tetap mengonfirmasi maksud pilihan. Tombol “Minta obat” hanya menyampaikan permintaan; sistem tidak memilih obat.')
p('Gambar membuktikan keberadaan antarmuka. Keterdengaran TTS, kesesuaian suara dengan pilihan, pencatatan ketika jaringan putus, dan waktu respons perlu diuji langsung. Papan tidak diposisikan sebagai pengganti panggilan perawat resmi.')

page('4.1 Ruangan dan tautan','Komunikasi pasien atau keluarga')
p('Tab Ruangan menyediakan label tempat tidur/ruangan dan tombol pembuatan tautan. Tautan dapat dibuka oleh pasien atau keluarga untuk mengirim pilihan kepada pemilik ruangan. Label sederhana membantu menghindari penulisan identitas lengkap pada papan.')
pair([
 ('papan komunikasi share link.png','Keadaan awal tab Ruangan sebelum ada ruangan.'),
 ('contoh pembuatan share link .png','Ruangan terbentuk dengan tindakan salin, bagikan, nonaktifkan, dan hapus.')
],4.7)
h('Hasil yang terlihat')
p('Dua gambar menunjukkan perubahan dari belum ada ruangan menjadi satu ruangan yang dapat dikelola. Keberadaan tombol tidak membuktikan bahwa tautan telah diterima atau dipakai berhasil pada perangkat kedua.')
h('Tindak lanjut pengujian')
p('Buka tautan dari perangkat lain, kirim satu pilihan fiktif, lalu periksa event pada akun pemilik. Setelah tautan dinonaktifkan, ulangi pengiriman dan pastikan ditolak. Catat pula perilaku jaringan putus dan apakah pengiriman ganda terjadi. Tautan tidak diperlakukan sebagai akses bebas ke riwayat pasien.')

page('4.2 Catatan Papan Komunikasi','Dari pilihan menjadi dokumentasi')
p('Tab Catatan menampilkan urutan kebutuhan dengan waktu, disertai tombol Salin untuk anamnesis dan Buka Dokumentasi. Jalur ini membantu perawat meninjau kembali pilihan yang diberikan pasien sebelum memasukkannya ke catatan.')
figure('catatan papan komunikasi.png','Catatan bertimestamp dari pilihan papan, dengan akses salin dan dokumentasi.',4.8)
h('Hasil pengamatan')
p('Gambar memperlihatkan beberapa pilihan seperti Ke toilet, Ubah posisi, Mual, Haus, dan Lapar. Catatan pilihan tetap merupakan laporan kebutuhan, bukan kesimpulan klinis. Ketika disalin ke anamnesis, informasi perlu diperiksa agar konteks waktu dan urutannya tidak hilang.')
p('Kode menyediakan penyimpanan lokal sebagai cadangan dan penyimpanan Supabase bagi pengguna login. Keberhasilan sinkronisasi seluruh event setelah jaringan kembali, pemisahan antar-akun, dan kesesuaian isi salinan perlu dibuktikan dengan skenario uji; screenshot saja belum mencakup kondisi itu.')

page('5. Anamnesis Terpandu','Tambahan mandiri dengan pertanyaan sistem')
p('Anamnesis Terpandu menyediakan pertanyaan yang dibacakan sistem melalui TTS. Pasien menjawab dengan suara, STT menghasilkan teks, lalu perawat dapat mengoreksi jawaban sebelum draf disusun. Implementasi ditemukan pada layar native Anamnesis; tidak ada tangkapan layar sesi lengkap modul ini dalam folder Progres 2.')
table(['Urutan','Topik pertanyaan pada implementasi'],[
 ('1','Keluhan utama yang dirasakan.'),('2','Waktu mulai atau durasi keluhan.'),('3','Lokasi keluhan dan beratnya pada skala 0–10.'),('4','Riwayat penyakit sebelumnya.'),('5','Alergi obat atau makanan.'),('6','Obat yang sedang rutin digunakan.'),('7','Keluhan lain yang ingin disampaikan.')
],[2,14.3])
h('Urutan penggunaan')
p('Mulai sesi dengan kode/inisial → dengarkan pertanyaan sistem → rekam jawaban → periksa hasil STT → koreksi atau lanjutkan → susun draf → periksa dan simpan → buka kembali melalui Dokumentasi & Riwayat.')
h('Contoh simulasi fiktif')
p('Sistem: “Sejak kapan keluhan itu mulai Anda rasakan?”\nPasien simulasi: “Sejak kemarin sore.”\nPerawat memeriksa apakah hasil teks tetap menyebut waktu yang sama sebelum jawaban dipakai dalam draf.')
box('Batas akademik dan bukti','Suara penanya adalah TTS sistem, bukan rekaman perawat. Gambar riwayat menunjukkan Anamnesis = 0. Laporan ini menyatakan implementasi tersedia, belum menyatakan alur lengkap telah berhasil diuji. Hasil dari modul ini tidak boleh disamakan begitu saja dengan percakapan alami perawat–pasien.')
p('Pengujian berikutnya mencakup TTS terdengar, izin mikrofon, jawaban STT, koreksi, kegagalan jaringan, penyimpanan, serta kesesuaian draf dan transkrip saat dibuka kembali.')

page('6. Dokumentasi & Riwayat','Satu tempat untuk meninjau sumber data')
p('Halaman dokumentasi menggabungkan pencarian dan filter dari Bot BPJS, Anamnesis, serta Papan. Perawat dapat menelusuri data berdasarkan sumber sebelum membuka hasil yang perlu diperiksa.')
figure('dokumentasi penuh semua bot.png','Riwayat gabungan: Semua 7, Bot BPJS 4, Anamnesis 0, dan Papan 3. Identitas disamarkan.',5.8)
p('Angka tersebut merupakan keadaan pada saat screenshot diambil, bukan ukuran kelengkapan penelitian. Status seperti “Siap dikirim” atau “Sudah dikirim” tidak berarti dokter sudah memverifikasi; verifikasi dilakukan di luar aplikasi. Tampilan daftar juga belum membuktikan isi transkrip dan draf telah diperiksa atau ekspor berhasil.')

page('7. Try and error OpenVPN','Kegagalan awal dan hasil setelah perbaikan')
p('Percobaan awal menampilkan “Belum terhubung” dengan pesan yang menyalahkan login. Catatan uji 18 Oktober menunjukkan server sebenarnya terjangkau; perbandingan log mengarahkan masalah pada negosiasi cipher klien.')
pair([
 ('try&error pada vpn.jpeg','Kegagalan awal dengan pesan login; alamat disamarkan.'),
 ('percobaan koneksi VPN data asli.jpeg','Status Terhubung setelah perbaikan; alamat disamarkan.')
],4.7)
table(['Percobaan / bukti','Hasil dan pelajaran'],[
 ('Isi ulang kredensial','Tetap gagal; bukan dasar untuk menyimpulkan password salah.'),
 ('Bandingkan log laptop dan HP','Server memilih AES-128-CBC, sedangkan daftar cipher klien tidak menyertakannya.'),
 ('Menambahkan ncp-disable','Ditolak mesin sebagai opsi tidak dikenal pada percobaan tersebut.'),
 ('Menambahkan data-ciphers dengan CBC','Koneksi berhasil; catatan mencatat Initialization Sequence Completed.')
],[6,10.3])
p('Akar masalah didukung oleh log dan perubahan hasil setelah perbaikan, bukan pembacaan konfigurasi server. Keberhasilan ini berlaku pada Wi-Fi yang diuji; jaringan seluler dan kestabilan jangka panjang belum dibuktikan. Pesan galat login yang menyesatkan masih menjadi tindak lanjut.')

page('8. Try and error Terminal SSH','Ketergantungan pada koneksi VPN')
p('Terminal gagal dengan Connection timed out ketika tunnel tidak aktif. Catatan sumber menjelaskan VPN sempat berhenti setelah aplikasi di-force-stop. Setelah VPN disambungkan kembali, Terminal berhasil login ke server asli sampai prompt.')
pair([
 ('try&error pada termina.jpeg','Timeout saat koneksi SSH belum berhasil; host/akun disamarkan.'),
 ('percobaan koneksi terminal data asli.jpeg','Shell berhasil terbuka; identitas/alamat disamarkan.')
],4.9)
h('Hasil dan batas percobaan')
p('Keberhasilan membuka prompt membuktikan jalur koneksi dan login pada kondisi tersebut. Catatan uji menyatakan tidak ada perintah yang dijalankan di server asli. Pesan banner server yang tampak adalah keluaran login, bukan hasil perintah pengujian.')
h('Pelajaran perbaikan')
p('Pastikan VPN aktif sebelum menghubungkan SSH ke server di belakang tunnel. Hindari force-stop di tengah uji yang membutuhkan VPN tetap hidup. Pisahkan kegagalan jangkauan, tunnel, dan autentikasi ketika mencatat galat; timeout sendiri bukan bukti password salah. Fitur Terminal SSH ini tidak sama dengan implementasi tunnel VPN berbasis SSH penuh.')

page('9. Panel Admin','Login dan dashboard operasional')
p('Panel Admin disediakan sebagai halaman terpisah untuk mengelola operasional sistem. Cakupan terbaru meliputi akun/profil umum, metadata database, maintenance, pengumuman, audit, dan status. Panel ini bukan antarmuka dokter.')
pair([
 ('login admin.png','Form login Panel Admin.'),
 ('dashboard admin.png','Dashboard hitungan data, grafik sesi, dan aktivitas admin.')
],4.9)
h('Perubahan dari rancangan awal')
p('Menu katalog Bot, Provider AI, dan Pengaturan generik telah dikeluarkan dari panel. Pilihan peran dokter/perawat juga dikeluarkan dari pengelolaan akun. Maintenance dan Pengumuman tetap dipertahankan. Perubahan ini tidak otomatis menghapus tabel konfigurasi yang sebelumnya sudah ada.')
p('HANDOFF 18 Oktober mencatat deploy admin/web di Fly dan uji desktop serta Chrome HP. Gambar dashboard menyajikan hitungan saat itu; hitungan tersebut bukan hasil evaluasi speech dan tidak membuktikan seluruh modul telah dipakai dengan sukses.')

page('9.1 Akun pengguna','Pengelolaan akun umum')
p('Menu Akun pengguna menampilkan status akun dan waktu terkait, menyediakan tombol Tambah pengguna dan Kelola. Identitas akun pada gambar laporan disamarkan; status serta struktur tabel dipertahankan.')
figure('menu akun pengguna.png','Pengelolaan akun umum tanpa pemilih peran; identitas disamarkan.',5.3)
h('Hasil yang terlihat')
p('Tabel berisi akun dengan status aktif atau belum konfirmasi. Keberadaan akun bukan bukti peran klinis, penggunaan oleh tenaga kesehatan, atau persetujuan dokter. Panel tidak memberikan akses isi klinis melalui halaman ini.')
p('Uji fixture yang tercatat memeriksa penolakan perubahan role serta pembuatan akun umum. Laporan Progres 2 tidak melakukan pembuatan, penghapusan, atau perubahan akun baru; gambar hanya digunakan untuk menjelaskan pengelolaan yang tersedia.')

page('9.2 Metadata sesi dan draf','Pemeriksaan data tanpa isi klinis')
p('Panel membedakan keberadaan record dari isi klinisnya. Daftar Sesi Bot BPJS menampilkan metadata sesi; daftar Dokumentasi LLM menampilkan metadata draf dan provider. Kolom transkrip maupun isi ringkasan tidak disajikan pada kedua halaman.')
pair([
 ('menu database log bot BPJS.png','Metadata sesi BPJS; ID disamarkan.'),
 ('menu database dokumentasi llm.png','Metadata dokumentasi LLM; ID disamarkan.')
],4.5)
h('Makna data pada gambar')
p('Status sesi dan tanggal membantu penelusuran operasional. Nama provider menjelaskan layanan yang tercatat membuat draf, bukan jaminan mutu hasilnya. Ada empat record pada masing-masing gambar; angka itu tidak menunjukkan WER, kelengkapan, atau factual consistency.')
h('Batas privasi')
p('Service key admin berada di server dan memiliki akses luas, sehingga pembatasan kolom pada server tetap diperlukan. Implementasi membatasi resource pasien menjadi metadata read-only. Laporan ini tidak mengubah skema atau RLS, dan tidak menyimpulkan keamanan menyeluruh hanya dari tampilan tabel.')

page('9.3 Segmen dan aktivitas pengguna','Jejak operasional yang dapat ditelusuri')
pair([
 ('menu database segmen transkrip.png','Metadata segmen: speaker, offset dan waktu; ID disamarkan.'),
 ('menu database log pengguna.png','Kategori/status aktivitas pengguna; ID disamarkan.')
],5.3)
h('Pengamatan pada metadata segmen')
p('Gambar menampilkan 11 segmen dan label speaker yang terlihat adalah “perawat”. Karena isi transkrip tidak ditampilkan, gambar ini tidak cukup untuk menilai apakah pembicara sudah tepat. Metadata ini juga bukan bukti bahwa diarization pada sesi lain berhasil atau gagal.')
h('Pengamatan pada log aktivitas')
p('Log memperlihatkan kategori seperti VPN dan AI serta badge informasi/sukses. Badge adalah catatan aplikasi, sehingga tetap perlu dikaitkan dengan bukti hasil nyata. Jumlah log tidak sama dengan jumlah percobaan yang berhasil; satu alur dapat menulis beberapa event.')
p('Penilaian diarization dan isi draf harus memakai transkrip, audio acuan, serta pemeriksaan terpisah oleh pihak yang berhak. Panel admin tidak digunakan sebagai jalan untuk membuka isi percakapan.')

page('9.4 Maintenance','Pengaturan pemeliharaan lapisan web')
p('Halaman Maintenance menyediakan pesan, jadwal mulai/selesai, dan pilihan tampilan. Fitur ini membantu pengelola memberi tahu pengguna ketika aplikasi web perlu dipelihara.')
figure('menu maintenance admin.png','Form Maintenance: pesan, jadwal, pilihan blokir/banner, dan tombol penerapan.',5.1)
h('Perilaku yang perlu dijelaskan')
p('Overlay maintenance diterapkan pada lapisan web. Implementasi mengecualikan tautan pasien/keluarga Papan Komunikasi. Layar native Bot BPJS dan Anamnesis tidak otomatis diblokir oleh overlay web, sehingga cakupannya tidak boleh dilaporkan sebagai penghentian seluruh fungsi aplikasi.')
p('Gambar menunjukkan formulir pengaturan, bukan bukti pengumuman maintenance telah diterima semua pengguna. Uji terpisah diperlukan untuk jadwal, kondisi sebelum/sesudah pemeliharaan, pengecualian tautan papan, serta pemulihan tampilan setelah maintenance berakhir.')

page('9.5 Pengumuman','Daftar dan formulir penambahan')
p('Pengumuman memiliki daftar dan formulir untuk judul, isi, jenis, serta rentang tayang. Pada bukti yang diberikan, daftar masih kosong dan gambar lainnya menunjukkan dialog penambahan.')
pair([
 ('menu pengumuman admin.png','Daftar pengumuman dalam keadaan belum ada data.'),
 ('contoh tambah pengumuman admin.png','Form tambah pengumuman sebelum pengisian/penyimpanan.')
],4.9)
h('Hasil dan tindak lanjut')
p('Keberadaan daftar dan formulir membuktikan antarmuka tersedia. Belum dapat disimpulkan bahwa satu pengumuman sudah disimpan, dijadwalkan, dan muncul pada perangkat pengguna dari kedua gambar ini.')
p('Pengujian berikutnya menggunakan isi uji fiktif: simpan pengumuman, periksa waktu tayang, buka aplikasi penerima, lalu nonaktifkan/hapus dan pastikan tampilan berubah sesuai. Pengujian tersebut perlu dicatat dengan waktu dan perangkat; tidak dijalankan sebagai bagian penyusunan laporan ini.')

page('9.6 Audit dan status layanan','Pemantauan panel')
pair([
 ('menu log admin.png','Audit tindakan admin, dengan ID target disamarkan.'),
 ('menu status admin.png','Status Supabase/web/relay dan keberadaan tabel panel.')
],5.1)
h('Audit log')
p('Audit mencatat tindakan administratif seperti login dan pengelolaan akun. Catatan tersebut mendukung penelusuran operasional, bukan bukti verifikasi diagnosis. Laporan ini tidak mengulangi tindakan yang terlihat pada log.')
h('Status & Setup')
p('Gambar memperlihatkan indikator OK untuk Supabase, web, dan relay, serta keberadaan app_settings, announcements, dan admin_audit_log. Angka respons pada halaman adalah satu pengamatan saat itu, bukan rata-rata performa atau jaminan uptime.')
p('Cuplikan komentar SQL pada screenshot masih menyebut katalog lama. Komentar historis itu tidak berarti menu katalog aktif kembali; cakupan panel terbaru mengikuti implementasi dan keputusan penyederhanaan yang dijelaskan pada Bagian 9.')

page('10. Evaluasi dan batas hasil','Membedakan bukti tampilan dari hasil ukur')
p('Progres 2 menyediakan dokumentasi perkembangan dan proses pemecahan masalah. Untuk menjawab penelitian, pengujian komparatif perlu menggunakan audio yang sama pada baseline STT saja dan proposed STT + LLM, disertai transkrip acuan manual.')
table(['Aspek','Bukti yang perlu dikumpulkan','Status pada laporan'],[
 ('WER','Jumlah substitusi, penghapusan, penyisipan dibanding kata acuan; WER = (S + D + I) / N.','Tidak ada angka WER baru dari screenshot.'),
 ('Kelengkapan','Daftar informasi acuan dan proporsi informasi yang tercatat.','Belum tersedia penilaian baru.'),
 ('Kesesuaian isi','Kecocokan data hasil dengan percakapan/acuan.','Perlu pemeriksaan isi, bukan hitungan record.'),
 ('Factual consistency','Catat fakta yang didukung, bertentangan, atau ditambahkan tanpa sumber.','Belum tersedia hasil review baru.'),
 ('Waktu proses','Titik mulai/akhir STT dan LLM; kondisi jaringan, model dan durasi audio.','Angka pada Status admin bukan waktu proses speech.'),
 ('Diarization/peran','Bandingkan label segmen dengan anotasi manual; pisahkan koreksi pengguna.','Label metadata saja tidak memadai.'),
 ('Papan Komunikasi','Pilihan → TTS → event → diterima pemilik; link nonaktif, offline, dan akses akun.','UI/catatan terlihat; uji lintas perangkat belum lengkap.'),
 ('Anamnesis Terpandu','Pertanyaan TTS, STT jawaban, koreksi, draf, simpan dan buka ulang.','Implementasi ada; bukti sesi lengkap belum tersedia.')
],[3.1,8.8,4.4])
box('Pemisahan baseline','Jika transkrip telah melewati relabel atau koreksi LLM, tandai tahap itu secara terpisah. Simpan keluaran STT sebelum pengolahan tambahan agar baseline STT saja tetap dapat dibandingkan secara jujur. Pertanyaan TTS pada Anamnesis jangan disebut rekaman perawat.')
p('Seluruh skenario baru memakai simulasi fiktif. Hasil draf harus dapat diperiksa terhadap sumbernya; tidak ada diagnosis atau saran obat otomatis yang dinilai sebagai keluaran sistem.')

page('11. Tindak lanjut dan penutup','Rencana setelah Progres 2')
table(['Prioritas','Langkah berikutnya','Bukti selesai'],[
 ('1','Membahas Papan dan Anamnesis yang ditambahkan mandiri dengan dosen.','Tanggapan/keputusan asli tercatat; jangan menganggap diam sebagai persetujuan.'),
 ('2','Menjalankan BPJS dari rekam sampai koreksi dan ekspor menggunakan skenario fiktif.','Audio, acuan, transkrip, draf, koreksi, dan ekspor dapat dicocokkan.'),
 ('3','Menjalankan satu sesi Anamnesis lengkap pada HP.','TTS/STT terdengar/tercatat, jawaban dikoreksi, draf tersimpan dan muncul di riwayat.'),
 ('4','Menguji papan di dua perangkat.','Event diterima pemilik, suara sesuai, link nonaktif ditolak, galat jaringan terlihat.'),
 ('5','Mengukur baseline/proposed dan mengulang skenario sulit.','WER, kelengkapan, kesesuaian, factual consistency, dan waktu tersedia.'),
 ('6','Memperbaiki penjelasan galat VPN dan menguji jaringan lain.','Pesan sesuai penyebab; seluler/stabilitas dicatat terpisah.'),
 ('7','Menguji penerapan pengumuman/maintenance secara terkendali.','Jadwal dan tampilan pada perangkat penerima sesuai; pengecualian papan bekerja.')
],[2,8.4,5.9])
h('Penutup laporan')
p('Perkembangan Progres 2 menunjukkan perluasan prototipe pada komunikasi pasien, dokumentasi gabungan, dan pengelolaan operasional. Percobaan VPN/Terminal menghasilkan pembelajaran yang jelas: pesan awal tidak selalu menunjukkan akar masalah; negosiasi cipher dan keberadaan tunnel perlu diperiksa melalui bukti.')
p('Papan Komunikasi dan Anamnesis Terpandu dilaporkan sebagai tambahan mandiri sambil menunggu jawaban dosen. Cakupan PA tetap berpusat pada dokumentasi percakapan perawat–pasien untuk membantu peninjauan dokter di luar aplikasi. Hasil yang tersedia menjadi dasar pengujian berikutnya, belum menjadi kesimpulan efektivitas klinis atau kelayakan produksi menyeluruh.')

page('Sumber penyusunan','Dokumen dan bukti internal')
p('Laporan ini memakai sumber lokal proyek. Tidak ada klaim bahwa referensi ilmiah lama telah diverifikasi ulang pada penyusunan ini. Angka dan hasil operasional mengikuti catatan sumber, sedangkan interpretasi gambar dibatasi pada hal yang terlihat.')
table(['Kode','Sumber','Pemakaian'],[
 ('S1','docs/word/referensi/Dokumentasi macam speech perawat dan pasien.docx','Acuan bentuk penjelasan dan daftar 9 komunikasi.'),
 ('S2','docs/word/pengajuan/Pengajuan-Pemilihan-Sistem-Speech.docx; ANALISIS-PEMILIHAN-SISTEM-SPEECH.md','Dua tambahan yang diajukan dan dasar prioritas.'),
 ('S3','Dokumentasi/Progres 1/','Latar tahap sebelumnya; 16 gambar, tidak digunakan sebagai gambar hasil Progres 2.'),
 ('S4','Dokumentasi/Progres 2/','21 gambar sumber seluruh ilustrasi hasil dalam laporan ini.'),
 ('S5','docs/hasil percobaan vpn dan terminal asli.md','Kronologi 18 Oktober, cipher, tunnel, keberhasilan dan batas.'),
 ('S6','docs/HANDOFF.md; docs/PANEL-ADMIN.md','Penyederhanaan, deployment dan catatan uji admin/web/HP.'),
 ('S7','docs/pengujian-minggu-ini/; docs/ujitest dan dokumentasi.md','Konteks uji lama: fixture, Docker, RLS, Flutter dan batasnya.'),
 ('S8','web/public/views/{bots,papan-komunikasi,bpjs-review}.js','Keberadaan modul, pilihan papan, ruangan, dan riwayat.'),
 ('S9','app/lib/screens/anamnesis_terpandu_screen.dart; app/lib/services/bpjs_service.dart','Pertanyaan TTS, jawaban STT dan alur draf; dibaca tanpa perubahan kode.'),
 ('S10','admin/lib/resources.js; admin/public/js/','Cakupan resource, metadata, akun dan halaman admin.'),
 ('S11','AGENTS.md; docs/KEPUTUSAN.md; PENGUJIAN-SPEECH-REALISTIS.md','Batas akademik, verifikasi di luar aplikasi, dan rencana evaluasi.')
],[1.2,9.5,5.6])
p('Gambar laporan menggunakan salinan yang disamarkan. Sumber asli, nama berkas, dan manifest pemetaan tersedia di folder proyek. Penyusunan: 19 Oktober 2026; tanggal percobaan mengikuti sumber, bukan tanggal penyusunan laporan.')

page('Daftar gambar dan penataan berkas','Lampiran penelusuran')
table(['Gambar','Nama berkas sumber di Progres 2'],[(str(n),name) for n,name,_ in figures],[2,14.3])
h('Susunan dokumentasi setelah dirapikan')
p('Word lama: docs/word/referensi, pengajuan, laporan, dan arsip. Non-Word: Markdown pada docs serta folder pengujian/aset; jalur dipertahankan agar rujukan kode dan aturan proyek tetap berlaku. Snapshot sistem lama: docs/arsip. Indeks lengkap: docs/README.md.')
p('Laporan final: Dokumentasi/Progres 2/progres 2 Try and error sistem.docx. Dokumen acuan “Dokumentasi macam speech perawat dan pasien.docx” tetap tersedia utuh di docs/word/referensi.')

assert len(figures)==21, len(figures)
assert len(pages)==23, len(pages)
doc.save(DEST)
(ASSETS/'report-structure.json').write_text(json.dumps({'pages_planned': pages,'figures':[{'number':n,'source':name,'caption':text} for n,name,text in figures]},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'report':str(DEST),'pages_planned':len(pages),'figures':len(figures),'bytes':DEST.stat().st_size}))
