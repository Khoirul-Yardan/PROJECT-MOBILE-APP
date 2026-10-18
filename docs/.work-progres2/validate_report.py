from pathlib import Path
import sys
import re
import json
import hashlib
import urllib.parse
import zipfile
import subprocess

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'docs/.work-speech/lib'))
from docx import Document
from PIL import Image,ImageDraw,ImageFont
import pymupdf as fitz
sys.stdout.reconfigure(encoding='utf-8')
DOCS = ROOT/'docs'
WORK = ROOT/'.work-docs-progres2'
REPORT = ROOT/'Dokumentasi/Progres 2/progres 2 Try and error sistem.docx'
findings={}
broken=[]
for path in DOCS.rglob('*.md'):
    if any(part.startswith('.work') for part in path.parts): continue
    for raw in re.findall(r'(?<!!)\[[^\]]*\]\((<[^>]*>|[^)]+)\)',path.read_text(encoding='utf-8-sig')):
        target=raw.strip('<>')
        if re.match(r'^[a-zA-Z]+:|^#',target): continue
        if not (path.parent/urllib.parse.unquote(target.split('#')[0])).exists():
            broken.append([str(path.relative_to(ROOT)),target])
findings['broken_markdown_links']=broken

moves={
 'Dokumentasi macam speech perawat dan pasien.docx':'referensi/Dokumentasi macam speech perawat dan pasien.docx',
 'Dokumentasi-Speech-Perawat-ke-Pasien.docx':'referensi/Dokumentasi-Speech-Perawat-ke-Pasien.docx',
 'Dokumentasi-Speech-Perawat-ke-Pasien-Revisi-Diskusi.docx':'referensi/Dokumentasi-Speech-Perawat-ke-Pasien-Revisi-Diskusi.docx',
 'Pengajuan-Pemilihan-Sistem-Speech.docx':'pengajuan/Pengajuan-Pemilihan-Sistem-Speech.docx',
 'arsip-speech/Dokumentasi-Speech-Perawat-ke-Pasien-v1.docx':'arsip/Dokumentasi-Speech-Perawat-ke-Pasien-v1.docx',
 'pengujian-minggu-ini/laporan mingguan.docx':'laporan/laporan mingguan.docx',
}
findings['original_word_preserved']={}
word_manifest_path=DOCS/'word/manifest.json'
word_manifest=json.loads(word_manifest_path.read_text(encoding='utf-8')) if word_manifest_path.exists() else {}
for old,new in moves.items():
    if old not in word_manifest:
        baseline=subprocess.run(['git','show','HEAD:docs/'+old],cwd=ROOT,capture_output=True,check=True).stdout
        word_manifest[old]={'new_path':'docs/word/'+new,'original_sha256':hashlib.sha256(baseline).hexdigest()}
    current=(DOCS/'word'/new).read_bytes()
    findings['original_word_preserved'][old]=word_manifest[old]['original_sha256']==hashlib.sha256(current).hexdigest()
word_manifest_path.write_text(json.dumps(word_manifest,ensure_ascii=False,indent=2),encoding='utf-8')
with zipfile.ZipFile(REPORT) as z:
    findings['zip_integrity_error']=z.testzip()
    findings['embedded_images']=len([n for n in z.namelist() if n.startswith('word/media/')])
    findings['external_image_relationship']=any('TargetMode="External"' in z.read(n).decode('utf-8') and 'relationships/image' in z.read(n).decode('utf-8') for n in z.namelist() if n.endswith('.rels'))
    embedded_hashes={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
    source_images=sorted(p for p in (ROOT/'Dokumentasi/Progres 2').iterdir() if p.suffix.lower() in ('.png','.jpg','.jpeg'))
    findings['embedded_images_identical_to_originals']={p.name:hashlib.sha256(p.read_bytes()).hexdigest() in embedded_hashes for p in source_images}
    findings['all_21_original_images_embedded']=len(source_images)==21 and all(findings['embedded_images_identical_to_originals'].values())
document=Document(REPORT)
findings['inline_images']=len(document.inline_shapes)
findings['tables']=len(document.tables)
findings['headings']=[p.text for p in document.paragraphs if p.style.name=='Heading 1']
manifest=json.loads((DOCS/'progres2-assets/manifest.json').read_text(encoding='utf-8'))
available=[row for row in manifest if (ROOT/'Dokumentasi/Progres 2'/row['source']).exists()]
findings['original_images_checked']=len(available)
findings['original_images_unchanged']=all(hashlib.sha256((ROOT/'Dokumentasi/Progres 2'/row['source']).read_bytes()).hexdigest()==row['source_sha256'] for row in available) if available else None
pdf=WORK/'report-preview.pdf'
if pdf.exists():
    pdfdoc=fitz.open(pdf)
    findings['rendered_pages']=len(pdfdoc)
    page_text=[]
    previews=WORK/'pages'; previews.mkdir(exist_ok=True)
    for i,page in enumerate(pdfdoc):
        lines=page.get_text().splitlines()
        page_text.append({'page':i+1,'characters':len(page.get_text()),'start':lines[:5],'end':lines[-4:]})
        page.get_pixmap(matrix=fitz.Matrix(1,1)).save(str(previews/f'{i+1:02}.png'))
    findings['page_text']=page_text
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',16)
    for batch in range((len(pdfdoc)+11)//12):
        sheet=Image.new('RGB',(1200,4*440),'#dddddd');draw=ImageDraw.Draw(sheet)
        for index in range(batch*12,min((batch+1)*12,len(pdfdoc))):
            im=Image.open(previews/f'{index+1:02}.png'); im.thumbnail((290,410))
            local=index-batch*12; x=(local%3)*400; y=(local//3)*440
            sheet.paste(im,(x,y));draw.text((x+300,y+15),str(index+1),fill='black',font=font)
        sheet.save(WORK/f'render-contact-{batch+1}.jpg')
(WORK/'validation.json').write_text(json.dumps(findings,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in findings.items() if k not in ('page_text','headings')},ensure_ascii=False,indent=2))
if 'page_text' in findings:
    for row in findings['page_text']:print(row['page'],row['characters'],' | '.join(row['start'][:4]))
