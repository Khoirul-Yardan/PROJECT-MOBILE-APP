from pathlib import Path
import sys, re, json
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT.parent / '.work-speech' / 'lib'))
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from bs4 import BeautifulSoup, NavigableString
import markdown

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2), Cm(2)
sec.header_distance, sec.footer_distance = Cm(.8), Cm(.8)
sec.different_first_page_header_footer = True
styles = doc.styles
for name in ['Normal', 'Body Text', 'List Bullet', 'List Number']:
    st = styles[name]
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor.from_string('253547')
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.12
    st.paragraph_format.widow_control = True
for name, size in [('Title', 34), ('Subtitle', 15), ('Heading 1', 21), ('Heading 2', 13), ('Heading 3', 11)]:
    st = styles[name]
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string('123D59')
    st.paragraph_format.space_before = Pt(14)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.keep_with_next = True
styles['Heading 1'].paragraph_format.page_break_before = True
for style_name in ['Source note', 'Code block']:
    styles.add_style(style_name, 1)
styles['Source note'].font.name = 'Calibri'
styles['Source note'].font.size = Pt(8)
styles['Source note'].font.color.rgb = RGBColor.from_string('687989')
styles['Source note'].paragraph_format.space_after = Pt(12)
styles['Code block'].font.name = 'Consolas'
styles['Code block'].font.size = Pt(8)
styles['Code block'].paragraph_format.space_after = Pt(0)
styles['Code block'].paragraph_format.line_spacing = 1.05
styles['Code block'].paragraph_format.left_indent = Cm(.2)
styles['Code block'].paragraph_format.right_indent = Cm(.2)

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd'); el.set(qn('w:fill'), fill); pr.append(el)

def field(p, instr, fallback=''):
    r = p.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    text = OxmlElement('w:instrText'); text.set(qn('xml:space'), 'preserve'); text.text = instr
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    r._r.append(begin); r._r.append(text); r._r.append(separate)
    p.add_run(fallback)
    p.add_run()._r.append(end)

header = sec.header.paragraphs[0]
header.text = 'AI HUB  /  LAPORAN PENGUJIAN MINGGUAN'
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor.from_string('687989')
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('17 Oktober 2026   |   ')
field(footer, ' PAGE ', '1')
footer.add_run(' / ')
field(footer, ' NUMPAGES ', '1')
for run in footer.runs:
    run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string('687989')

doc.core_properties.title = 'Laporan Mingguan — Pengujian AI Hub'
doc.core_properties.subject = 'Kompilasi laporan pengujian 17 Oktober 2026'
doc.core_properties.author = 'Tim Proyek AI Hub'
doc.core_properties.keywords = 'AI Hub, pengujian, web, VPN, RLS, Flutter'

p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(65)
r = p.add_run('AI HUB'); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor.from_string('14858A')
doc.add_paragraph('LAPORAN\nMINGGUAN', 'Title')
doc.add_paragraph('Pengujian aplikasi dan infrastruktur', 'Subtitle')
p = doc.add_paragraph('17 Oktober 2026'); p.paragraph_format.space_before = Pt(16)
p.runs[0].font.size = Pt(13)
doc.add_paragraph('Proyek Akhir D3 PENS Lamongan')
p = doc.add_paragraph('Web & Relay  •  VPN  •  Database & RLS\nFlutter  •  Pengujian HP / Emulator')
p.paragraph_format.space_before = Pt(22)
p.paragraph_format.space_after = Pt(30)
doc.add_paragraph('Tentang dokumen', 'Heading 2')
doc.add_paragraph('Dokumen ini menggabungkan seluruh isi enam berkas Markdown dalam folder pengujian-minggu-ini. Hasil, batasan, temuan, catatan transparansi, dan langkah pengujian dipertahankan. Tabel yang panjang disajikan sebagai rincian per tes agar mudah dibaca.')
doc.add_paragraph('Tanggal laporan mengikuti sumber. Penyusunan dokumen ini tidak mencakup pelaksanaan ulang pengujian.', 'Source note')

doc.add_paragraph('Daftar isi', 'Heading 1')
titles = ['1. Ringkasan pengujian', '2. Web dan Relay', '3. VPN — server tiruan di Docker', '4. Database dan RLS', '5. Flutter — pengujian di host', '6. Pengujian HP atau emulator']
for index, title in enumerate(titles, 1):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    link = OxmlElement('w:hyperlink'); link.set(qn('w:anchor'), 'bagian_' + str(index))
    run = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '123D59'); pr.append(color)
    run.append(pr)
    tx = OxmlElement('w:t'); tx.text = title; run.append(tx)
    link.append(run); p._p.append(link)
doc.add_paragraph('Panduan membaca', 'Heading 2')
doc.add_paragraph('Bagian 1 memuat ringkasan dan cara mengulang pengujian. Bagian 2–5 menyajikan rincian hasil per area. Bagian 6 merupakan lembar tindak lanjut untuk pengujian perangkat; kolom hasil pada sumber masih kosong dan disediakan sebagai ruang isian.')
doc.add_paragraph('Daftar sumber', 'Heading 2')
sources = ['README.md', '01-web-dan-relay.md', '02-vpn-docker.md', '03-database-rls.md', '04-flutter.md', '05-perlu-hp-emulator.md']
for s in sources:
    doc.add_paragraph(s, 'List Bullet')

def inline(p, node, bold=False, italic=False, code=False):
    if isinstance(node, NavigableString):
        r = p.add_run(str(node)); r.bold = bold; r.italic = italic
        if code:
            r.font.name = 'Consolas'; r.font.size = Pt(9)
        return
    if node.name == 'br':
        p.add_run().add_break(); return
    for child in node.children:
        inline(p, child, bold or node.name in ['strong','b'], italic or node.name in ['em','i'], code or node.name == 'code')

def table(node, source):
    rows = node.find_all('tr')
    heads = rows[0].find_all(['th','td'])
    # Long narrative tables become cards; all original cell content is retained.
    if source in ['01-web-dan-relay.md', '05-perlu-hp-emulator.md']:
        for row in rows[1:]:
            cells = row.find_all(['td','th'])
            device = source.startswith('05')
            title = (cells[0].get_text() + '. ' + cells[1].get_text()) if device else cells[0].get_text()
            doc.add_paragraph(title, 'Heading 3')
            start = 2 if device else 1
            for j in range(start, len(cells)):
                p = doc.add_paragraph()
                p.add_run(heads[j].get_text() + ': ').bold = True
                if cells[j].get_text().strip(): inline(p, cells[j])
                else: p.add_run('________________________________________________' if device else '—')
        return
    t = doc.add_table(rows=1, cols=len(heads))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if len(heads) == 6: widths = [.6, 3, 4.4, 5.8, 1.1, 2.1]
    elif len(heads) == 2: widths = [5.2, 11.8]
    elif len(heads) == 3: widths = [4.5, 5.7, 6.8]
    else: widths = [17 / len(heads)] * len(heads)
    # The summary has five columns.
    if len(heads)==5: widths=[.6,3.4,5,6.7,1.3]
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    for i,row in enumerate(rows):
        dest = t.rows[0] if i == 0 else t.add_row()
        cant = OxmlElement('w:cantSplit'); dest._tr.get_or_add_trPr().append(cant)
        if i == 0:
            repeat=OxmlElement('w:tblHeader'); dest._tr.get_or_add_trPr().append(repeat)
        for j,cell in enumerate(row.find_all(['th','td'])):
            target = dest.cells[j]; target.width = Cm(widths[j])
            target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            inline(target.paragraphs[0],cell)
            for p in target.paragraphs:
                p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4)
                p.paragraph_format.line_spacing=1.04
                for r in p.runs:
                    r.font.size=Pt(8.5 if len(heads)>3 else 9)
                    if i == 0: r.bold=True; r.font.color.rgb=RGBColor(255,255,255)
            shade(target, '123D59' if i==0 else ('F0F5F8' if i%2 else 'FFFFFF'))
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

for index,(source,title) in enumerate(zip(sources,titles),1):
    p = doc.add_paragraph(title,'Heading 1')
    start = OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(index)); start.set(qn('w:name'),'bagian_'+str(index))
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(index))
    p._p.insert(0,start); p._p.append(end)
    doc.add_paragraph('Sumber: '+source,'Source note')
    soup = BeautifulSoup(markdown.markdown((ROOT/source).read_text(encoding='utf-8-sig'), extensions=['tables','fenced_code']), 'html.parser')
    for el in soup.children:
        if isinstance(el,NavigableString): continue
        if el.name=='h1':
            p=doc.add_paragraph(el.get_text(),'Source note')
        elif el.name in ['h2','h3','h4']:
            doc.add_paragraph(el.get_text(),'Heading 2' if el.name=='h2' else 'Heading 3')
        elif el.name=='p': inline(doc.add_paragraph(),el)
        elif el.name in ['ul','ol']:
            for idx,li in enumerate(el.find_all('li',recursive=False),1):
                p=doc.add_paragraph(style='List Bullet' if el.name=='ul' else 'Normal')
                if el.name=='ol': p.add_run(str(idx)+'. ')
                inline(p,li)
        elif el.name=='table': table(el,source)
        elif el.name=='pre':
            for line in el.get_text().rstrip('\n').split('\n'):
                p=doc.add_paragraph(line,'Code block')
                sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'F0F5F8'); p._p.get_or_add_pPr().append(sh)
        else: raise ValueError('Unsupported HTML: '+el.name)

doc.save(ROOT/'laporan mingguan.docx')

# Verify every source text node survives conversion, allowing only heading layout changes.
alltext='\n'.join(doc.element.body.itertext())
missing=[]
for source in sources:
    soup=BeautifulSoup(markdown.markdown((ROOT/source).read_text(encoding='utf-8-sig'),extensions=['tables','fenced_code']),'html.parser')
    for node in soup.find_all(string=True):
        val=str(node).strip()
        if val and val not in alltext:
            # Code blocks are intentionally split into individual paragraphs.
            if not all(line.strip() in alltext for line in val.splitlines() if line.strip()): missing.append([source,val])
assert not missing, missing
print(json.dumps({'file':str(ROOT/'laporan mingguan.docx'),'sources':len(sources),'source_text_check':'PASS','bytes':(ROOT/'laporan mingguan.docx').stat().st_size}))
