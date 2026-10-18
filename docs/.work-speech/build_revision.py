import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent / 'lib'))
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import json

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / 'speech-assets' / 'bahan-diskusi'
ASSETS.mkdir(exist_ok=True)
INK = '#173448'; TEAL = '#117F83'; LIGHT = '#EAF5F3'; GOLD = '#D99637'
FONT = 'C:/Windows/Fonts/arial.ttf'
BOLD = 'C:/Windows/Fonts/arialbd.ttf'
def font(size, bold=False): return ImageFont.truetype(BOLD if bold else FONT, size)
def txt(d, xy, text, size=30, color=INK, bold=False):
    d.multiline_text(xy, text, font=font(size, bold), fill=color, spacing=12)
def person(d,x,y,nurse=False):
    color=TEAL if nurse else '#5378A2'
    d.ellipse((x-36,y,x+36,y+72),fill='#E6BA98',outline=INK,width=3)
    d.rounded_rectangle((x-65,y+84,x+65,y+205),radius=32,fill=color)
    d.line((x-30,y+200,x-30,y+250),fill=INK,width=17)
    d.line((x+30,y+200,x+30,y+250),fill=INK,width=17)
    d.ellipse((x-15,y+29,x-10,y+34),fill=INK);d.ellipse((x+10,y+29,x+15,y+34),fill=INK)
    d.arc((x-12,y+35,x+12,y+53),0,180,fill=INK,width=2)
    if nurse:
        d.rounded_rectangle((x+12,y+110,x+43,y+144),radius=4,fill='white')
        d.line((x+18,y+122,x+37,y+122),fill=TEAL,width=3)
    txt(d,(x-62,y+272),'PERAWAT' if nurse else 'PASIEN',22,bold=True)
def bubble(d,box,text,sz=30):
    d.rounded_rectangle(box,radius=20,fill='white',outline='#A6C8CB',width=3)
    txt(d,(box[0]+24,box[1]+22),text,sz)
def illustration(n):
    im=Image.new('RGB',(1400,510),'#F4F8FA');d=ImageDraw.Draw(im)
    if n==1:
        person(d,135,140,True);person(d,1245,140)
        bubble(d,(260,65,1115,190),'“Apa yang paling mengganggu Ibu?”',35)
        bubble(d,(260,225,1115,350),'“Saya merasa mual sejak kemarin.”',35)
        txt(d,(300,400),'Bertanya • mendengarkan • mengonfirmasi',32,TEAL,True)
    elif n==2:
        person(d,135,145,True);person(d,370,145)
        txt(d,(62,50),'Percakapan perawat–pasien',30,bold=True)
        d.rounded_rectangle((570,55,1330,450),radius=20,fill='white',outline='#B5C8D2',width=3)
        txt(d,(615,85),'CONTOH TEKS HASIL TRANSKRIP',27,TEAL,True)
        txt(d,(615,155),'Perawat: Ada muntah?\nPasien: Mual, tetapi tidak muntah.',31)
        d.line((615,290,1285,290),fill='#C6D6DF',width=3)
        txt(d,(615,325),'Perawat memeriksa kembali\nkecocokan teks dengan ucapan.',29)
    elif n==3:
        person(d,135,150,True);person(d,1245,150)
        d.rounded_rectangle((300,65,1080,420),radius=25,fill=INK)
        d.rounded_rectangle((320,85,1060,400),radius=15,fill='white')
        txt(d,(360,120),'UCAPAN YANG DIBACA PASIEN',27,TEAL,True)
        txt(d,(360,190),'“Saya akan memeriksa\ntekanan darah.”',44,bold=True)
        txt(d,(360,325),'Contoh ilustrasi caption',26)
    elif n==5:
        person(d,135,150);person(d,1245,150,True)
        d.rounded_rectangle((300,55,1080,420),radius=20,fill='white',outline='#B5C8D2',width=3)
        txt(d,(345,88),'PASIEN MENULIS PESANNYA',27,TEAL,True)
        txt(d,(345,155),'“Tolong bantu ubah posisi saya.”',34,bold=True)
        for row in range(3):
            for col in range(10):
                x=345+col*66;y=235+row*43
                d.rounded_rectangle((x,y,x+53,y+32),radius=5,fill='#E2EBF0')
        txt(d,(345,375),'Pesan dibacakan menjadi suara',28,TEAL)
    elif n==7:
        txt(d,(60,35),'PASIEN MENUNJUKKAN INTENSITAS NYERI',30,TEAL,True)
        txt(d,(60,98),'Contoh jawaban pasien: “Nyeri saya 4 dari 10.”',35,bold=True)
        for i in range(11):
            x=65+i*115
            d.rounded_rectangle((x,205,x+98,310),radius=15,fill=TEAL if i==4 else 'white',outline=TEAL,width=3)
            txt(d,(x+(30 if i<10 else 18),235),str(i),40,'white' if i==4 else INK,True)
        txt(d,(65,345),'0 = tidak nyeri',28)
        txt(d,(835,345),'10 = nyeri terburuk\nyang dapat dibayangkan',28)
        txt(d,(65,430),'Lokasi dan keluhan lain dijelaskan secara terpisah.',28,TEAL)
    elif n==8:
        person(d,135,155,True);person(d,1245,155)
        bubble(d,(270,45,1110,180),'Perawat menjelaskan informasi\ndengan bahasa yang mudah dipahami.',30)
        bubble(d,(270,210,1110,345),'Pasien menjelaskan kembali\ndengan kata-katanya sendiri.',30)
        txt(d,(300,410),'Perawat memperjelas bagian yang belum dipahami.',27,TEAL,True)
    p=ASSETS/f'ilustrasi-{n:02}.png';im.save(p);return p

sources=[
('Siti Fatimah / Kementerian Kesehatan RI (2022). Komunikasi Terapeutik.','https://keslan.kemkes.go.id/view_artikel/1344/komunikasi-terapeutik','Artikel edukasi profesi; dasar komunikasi yang berpusat pada pasien. Digunakan pada nomor 1.'),
('The Learning Disability Partnership Board in Surrey. The Hospital Communication Book. Salinan pada situs UHS NHS.','https://www.uhs.nhs.uk/Media/UHS-website-2019/Docs/For-patients/The-Hospital-Communication-book.pdf','Buku alat bantu komunikasi. Bahasa sederhana: hlm. 3; simbol: hlm. 8–14; bagian tubuh: hlm. 19–21. Digunakan pada nomor 1, 4, dan 7.'),
('NHS England (2025; diperbarui 2026). Guidance on the use of AI-enabled ambient scribing products in health and care settings.','https://www.england.nhs.uk/long-read/guidance-on-the-use-of-ai-enabled-ambient-scribing-products-in-health-and-care-settings/','Panduan layanan kesehatan Inggris; penjelasan transkripsi, draf catatan, dan pemeriksaan hasil. Digunakan pada nomor 2.'),
('Google. Use Live Transcribe. Android Accessibility Help.','https://support.google.com/accessibility/android/answer/9158064?hl=en','Dokumentasi produk untuk ucapan menjadi teks. Digunakan pada nomor 3.'),
('Norfolk and Norwich University Hospitals NHS Foundation Trust. Alternative and Augmentative Communication.','https://www.nnuh.nhs.uk/departments/speech-and-language-therapy/communication/alternative-and-augmentative-communication/','Penjelasan alat AAC sederhana dan elektronik. Digunakan pada nomor 4 dan 5.'),
('American Speech-Language-Hearing Association (ASHA). Communication Access Resources: Communication Aids.','https://www.asha.org/practice/communication-access/communication-aids/','Sumber organisasi profesi tentang alat komunikasi dan respons Ya/Tidak dengan bantuan pendamping. Digunakan pada nomor 5 dan 6.'),
('ASHA. Yes/No/I Don’t Know Badge Cards.','https://www.asha.org/siteassets/practice-management/yes-no-dont-know-sheet.pdf','Kartu komunikasi; cuplikan pada gambar 6 berasal dari halaman 1.'),
('Cambridge University Hospitals NHS Foundation Trust. Understanding how to manage your pain.','https://www.cuh.nhs.uk/our-services/fit-4-surgery/your-stay-in-hospital/understanding-how-to-manage-your-pain/','Informasi pasien tentang pelaporan nyeri 0–10 dan kesesuaian cara penilaian. Digunakan pada nomor 7.'),
('Agency for Healthcare Research and Quality (AHRQ, 2023). TeamSTEPPS: Tool—Teach-Back.','https://www.ahrq.gov/teamstepps-program/curriculum/communication/tools/teachback.html','Materi pendidikan komunikasi untuk memeriksa pemahaman pasien. Digunakan pada nomor 8.'),
('Tobii Dynavox. TD Pilot.','https://us.tobiidynavox.com/pages/td-pilot','Dokumentasi produsen tentang perangkat komunikasi dengan akses tatapan. Digunakan pada nomor 9.'),
('Tobii Dynavox. TD Pilot Product Information Sheet.','https://downloads.tobiidynavox.com/Devices/TD_Pilot/TD_PI-Sheet_TD-Pilot_US.pdf','Lembar informasi produk; foto pada gambar 9 berasal dari halaman 1.')]

items=[
dict(title='Komunikasi terapeutik dan wawancara',tag='KOMUNIKASI LISAN',summary='Menggali keluhan dan membangun hubungan saling percaya.',
 explanation='Percakapan yang direncanakan, berpusat pada pasien, dan dilakukan dengan empati. Perawat mendengarkan keluhan, memberi kesempatan bercerita, lalu meminta penjelasan tambahan. Wawancara dapat dipandu topik agar informasi lebih terarah. [1]',
 nurse='Mengajukan pertanyaan yang mudah dipahami, mendengarkan, dan mengonfirmasi maksud pasien.',patient='Menceritakan keluhan, perasaan, atau kebutuhan dengan kata-katanya sendiri.',
 example='Perawat: “Apa yang paling mengganggu Ibu saat ini?”\nPasien: “Saya mual sejak kemarin sore.”\nPerawat: “Sejak kemarin sore, ya. Apa lagi yang Ibu rasakan?”',
 note='Gunakan kalimat pendek dan beri waktu untuk menjawab. Pertanyaan dalam contoh ini merupakan ilustrasi percakapan, bukan daftar pengkajian baku. [2]',refs=[1,2]),
dict(title='Ucapan menjadi teks untuk dokumentasi',tag='SPEECH-TO-TEXT / STT',summary='Membantu mengubah percakapan menjadi catatan tertulis.',
 explanation='Speech-to-text mengubah ucapan menjadi transkrip. Pada ambient scribing, percakapan dapat diolah lebih lanjut menjadi draf catatan. Transkrip dan ringkasan merupakan keluaran yang berbeda; keduanya perlu diperiksa oleh tenaga kesehatan. [3]',
 nurse='Melakukan wawancara dan memeriksa apakah hasil tulisan sesuai dengan ucapan.',patient='Menyampaikan keluhan atau riwayat secara lisan dan membantu mengklarifikasi informasi yang kurang jelas.',
 example='Perawat: “Apakah ada muntah?”\nPasien: “Mual, tetapi tidak muntah.”\nTeks yang diperiksa: pasien mengeluh mual dan menyatakan tidak muntah.',
 note='Dapat digunakan untuk membantu dokumentasi percakapan ketika perekaman diizinkan. Kesalahan kata, angka, atau negasi dapat mengubah makna; hasil otomatis tetap perlu dikoreksi. [3]',refs=[3]),
dict(title='Ucapan menjadi teks untuk dibaca pasien',tag='CAPTION / TRANSKRIP LANGSUNG',summary='Membantu pasien membaca pesan yang diucapkan perawat.',
 explanation='Ucapan perawat ditampilkan sebagai tulisan saat percakapan berlangsung. Contoh alat yang mendokumentasikan fungsi ini adalah Google Live Transcribe. Ukuran teks dapat disesuaikan dan balasan dapat diketik. [4]',
 nurse='Berbicara dengan jelas serta memastikan tulisan yang muncul sesuai dengan pesan.',patient='Membaca pesan, meminta pengulangan, lalu menjawab dengan suara atau tulisan sesuai pilihannya.',
 example='Perawat: “Saya akan memeriksa tekanan darah.”\nPasien membaca teks tersebut.\nPasien: “Sebentar, saya ingin duduk lebih nyaman.”',
 note='Berguna saat pasien memilih komunikasi tertulis. Kesesuaian bergantung pada keterbacaan, bahasa, dan ketepatan transkripsi. Dokumentasi produk menjelaskan fungsi alat, bukan jaminan ketepatan istilah klinis. [4]',refs=[4]),
dict(title='Papan gambar atau simbol komunikasi',tag='KOMUNIKASI ALTERNATIF / AAC',summary='Menyampaikan kebutuhan dengan menunjuk gambar atau kata.',
 explanation='Pasien memilih simbol yang mewakili pesan, misalnya merasa panas, sakit kepala, atau pusing. Papan dapat berupa lembar cetak atau alat elektronik; pada perangkat bersuara, pilihan dapat dibacakan. [2][5]',
 nurse='Menawarkan pilihan yang relevan dan memastikan makna gambar yang dipilih.',patient='Menunjuk gambar atau kata untuk menyampaikan kebutuhan tanpa harus mengucapkan kalimat.',
 example='Pasien menunjuk simbol “Headache”.\nPerawat: “Apakah maksudnya kepala terasa sakit?”\nPasien mengonfirmasi dengan cara komunikasi yang disepakati.',
 note='Sesuai ketika pasien dapat menggunakan gambar atau kata sebagai bantuan komunikasi. Pemahaman terhadap simbol perlu dipastikan; tidak semua orang memaknai gambar dengan cara yang sama. [2]',refs=[2,5],image='rujukan-nhs-simbol.png',caption='Cuplikan simbol gejala dari The Hospital Communication Book, hlm. 14. Sumber: [2]; simbol dikreditkan buku kepada Widgit. Label asli berbahasa Inggris dipertahankan.'),
dict(title='Mengetik atau memilih huruf menjadi suara',tag='TEXT-TO-SPEECH / TTS',summary='Menyampaikan pesan bebas melalui tulisan yang dibacakan.',
 explanation='Pasien menyusun pesan tertulis, kemudian alat elektronik membacakannya. Cara ini termasuk AAC dengan keluaran suara. Papan huruf dan alat komunikasi elektronik merupakan contoh sarana AAC yang dijelaskan layanan terapi wicara. [5][6]',
 nurse='Memberi waktu untuk menyusun pesan, mendengarkan hasilnya, dan mengonfirmasi maksudnya.',patient='Mengetik atau memilih huruf untuk menyampaikan kalimat yang diinginkan.',
 example='Pasien mengetik: “Tolong bantu ubah posisi saya.”\nAlat membacakan kalimat tersebut.\nPerawat: “Bapak ingin dibantu mengubah posisi, benar?”',
 note='Bisa dibahas untuk pasien yang memahami tulisan tetapi sulit menghasilkan ucapan. Papan huruf cetak dapat membantu mengeja dengan pendamping; papan cetak sendiri tidak menghasilkan suara elektronik. [5][6]',refs=[5,6]),
dict(title='Jawaban Ya, Tidak, dan Tidak Tahu',tag='RESPONS SEDERHANA DENGAN PENDAMPING',summary='Membantu pasien memberikan jawaban singkat yang jelas.',
 explanation='Kartu pilihan memungkinkan pasien menjawab tanpa berbicara. Bila menunjuk sulit, pendamping dapat menyebutkan pilihan satu per satu dan menunggu sinyal yang telah disepakati. ASHA menjelaskan penggunaan sistem Ya/Tidak serta menyediakan kartu tiga respons. [6][7]',
 nurse='Mengajukan satu pertanyaan setiap kali dan memastikan cara pasien menunjukkan pilihan.',patient='Memilih Ya, Tidak, atau Tidak Tahu melalui penunjukan atau sinyal yang dapat digunakan.',
 example='Perawat: “Apakah Ibu merasa mual?”\nPasien menunjuk “Ya”.\nPerawat mengonfirmasi jawaban sebelum melanjutkan pertanyaan.',
 note='“Tidak tahu” berbeda dari “tidak”. Tidak adanya respons juga tidak dapat dianggap sebagai jawaban tidak. Kemampuan melihat, mendengar, dan memberi sinyal perlu diperhatikan. [6]',refs=[6,7],image='rujukan-asha-kartu.png',caption='Kartu ASHA: Yes / No / I Don’t Know, hlm. 1. Sumber gambar: [7]. Arti dalam bahasa Indonesia: Ya / Tidak / Tidak Tahu.'),
dict(title='Komunikasi lokasi dan intensitas nyeri',tag='PELAPORAN NYERI OLEH PASIEN',summary='Membantu menjelaskan bagian yang nyeri dan tingkat nyerinya.',
 explanation='Pasien menjelaskan lokasi nyeri dan seberapa berat nyeri yang dirasakan. Salah satu cara ialah skala angka 0–10: 0 berarti tidak nyeri, sedangkan 10 berarti nyeri terburuk yang dapat dibayangkan. Gambar tubuh dapat membantu menunjukkan lokasi. [2][8]',
 nurse='Menanyakan lokasi dan intensitas, lalu memastikan jawaban mencerminkan laporan pasien.',patient='Menyebutkan atau menunjuk lokasi tubuh serta memilih angka yang mewakili nyerinya.',
 example='Perawat: “Di bagian mana terasa nyeri? Jika 0 tidak nyeri dan 10 nyeri terburuk yang dapat dibayangkan, berapa nilainya?”\nPasien: “Di perut bagian bawah, nilainya 4.”',
 note='Angka adalah laporan pasien, bukan diagnosis atau penentuan obat otomatis. Skala angka tidak sesuai untuk semua orang; cara lain dibahas bersama tenaga kesehatan bila pasien sulit memilih angka. [8]',refs=[2,8]),
dict(title='Edukasi dengan metode teach-back',tag='MEMERIKSA PEMAHAMAN PASIEN',summary='Memastikan penjelasan perawat dipahami dengan benar.',
 explanation='Setelah menerima penjelasan, pasien diminta menyampaikan kembali informasi menggunakan kata-katanya sendiri. Perawat mendengarkan, memperjelas bagian yang belum dipahami, lalu mengecek kembali. Fokusnya ialah kejelasan penjelasan yang diberikan. [9]',
 nurse='Menjelaskan dengan bahasa sederhana dan meminta pasien menceritakan kembali inti informasi.',patient='Menjelaskan apa yang dipahami serta menyampaikan bagian yang masih membingungkan.',
 example='Perawat: “Supaya saya tahu penjelasan saya sudah jelas, boleh ceritakan kembali apa yang perlu dilakukan setelah pulang sesuai petunjuk tadi?”\nPasien menjelaskan dengan kata-katanya sendiri; perawat memperjelas bagian yang belum sesuai.',
 note='Dapat digunakan saat edukasi perawatan atau penjelasan tindak lanjut. Pertanyaan “Sudah paham?” saja tidak menunjukkan apa yang benar-benar dipahami pasien. [9]',refs=[9]),
dict(title='Komunikasi dengan pelacakan mata',tag='AAC DENGAN AKSES TATAPAN',summary='Memilih pesan dengan tatapan ketika akses tangan sulit.',
 explanation='Pelacakan mata memungkinkan pengguna memilih huruf atau simbol melalui arah tatapan pada perangkat yang mendukungnya. Pesan yang disusun dapat dibacakan. TD Pilot dari Tobii Dynavox merupakan salah satu contoh perangkat komunikasi dengan akses tersebut. [10][11]',
 nurse='Membantu menempatkan alat, memberi waktu untuk memilih pesan, dan mengonfirmasi hasil komunikasi.',patient='Memilih pesan melalui tatapan sesuai kemampuan dan pengaturan aksesnya.',
 example='Pasien memilih pesan “Saya ingin istirahat” melalui tatapan.\nPerangkat membacakan pesan.\nPerawat memastikan bahwa pesan tersebut sesuai dengan maksud pasien.',
 note='Merupakan cara mengakses alat komunikasi, bukan pengenalan ucapan. Kesesuaian perangkat dan pengaturan perlu dinilai secara individual. Dokumentasi produsen menjelaskan produk, bukan bukti bahwa alat sesuai untuk setiap pasien. [10][11]',refs=[10,11],image='rujukan-td-pilot.png',caption='Contoh bentuk perangkat TD Pilot. Foto dari TD Pilot Product Information Sheet, hlm. 1. Sumber: [11], © Tobii Dynavox; ditampilkan sebagai contoh alat komunikasi.')
]

doc=Document();s=doc.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=Cm(1.9);s.bottom_margin=Cm(1.8);s.left_margin=Cm(2);s.right_margin=Cm(2)
s.header_distance=Cm(.8);s.footer_distance=Cm(.8)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor.from_string('173448')
normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.12
for name,size in [('Title',27),('Heading 1',20),('Heading 2',12)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor.from_string('117F83')
    st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(8)
doc.styles['Caption'].font.name='Arial';doc.styles['Caption'].font.size=Pt(8.5)
doc.styles['Caption'].font.color.rgb=RGBColor.from_string('526575')
hp=s.header.paragraphs[0];hp.text='SPEECH & KOMUNIKASI PERAWAT–PASIEN  |  BAHAN DISKUSI';hp.runs[0].font.size=Pt(8);hp.runs[0].font.color.rgb=RGBColor.from_string('597482')
fp=s.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('Bahan diskusi dengan dosen  •  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
for r in fp.runs:r.font.size=Pt(8)
doc.core_properties.title='Daftar Speech dan Komunikasi Perawat–Pasien'
doc.core_properties.subject='Daftar, penjelasan, contoh interaksi, gambar, dan sumber untuk diskusi dengan dosen'
doc.core_properties.author=''
def para(t,style=None): return doc.add_paragraph(t,style)
def label(l,t):
    p=para('');p.add_run(l+' ').bold=True;p.add_run(t);return p
def link(p,text,url):
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'117F83');pr.append(c)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,h in zip(t.rows[0].cells,headers):
        c.text=h;sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'117F83');c._tc.get_or_add_tcPr().append(sh)
        for r in c.paragraphs[0].runs:r.font.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for c,text in zip(cells,row):
            c.text=text
            if i%2==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EFF5F7');c._tc.get_or_add_tcPr().append(sh)
    for row in t.rows:
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if widths:
            for c,w in zip(row.cells,widths):c.width=Cm(w)
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(6);p.paragraph_format.space_before=Pt(6)
                for r in p.runs:r.font.size=Pt(9)
    return t

para('BAHAN DISKUSI DENGAN DOSEN','Subtitle')
doc.add_heading('Daftar Speech dan\nKomunikasi Perawat–Pasien',0)
para('Jenis penggunaan • Penjelasan • Gambar • Sumber','Subtitle')
para('Dokumen ini merangkum sembilan bentuk komunikasi yang dapat dibahas bersama dosen. Setiap jenis dijelaskan melalui peran perawat dan pasien, contoh interaksi, gambar pendamping, serta rujukan yang dapat dibuka.')
para('Istilah speech di sini mencakup ucapan lisan, ucapan menjadi teks (STT), dan teks menjadi suara (TTS). Papan gambar, kartu jawaban, dan akses tatapan juga disertakan sebagai komunikasi augmentatif dan alternatif (AAC), yaitu cara membantu atau menggantikan ucapan. [5][6]')
doc.add_heading('Daftar yang dibahas',2)
table(['No.','Jenis komunikasi','Kegunaan utama'],[(str(i),x['title'],x['summary']) for i,x in enumerate(items,1)],[1,7.7,8.3])
para('Nomor 1–9 adalah urutan pembahasan. Pemilihan cara komunikasi mengikuti kebutuhan dan kemampuan pasien. Seluruh contoh percakapan dibuat untuk penjelasan dan menggunakan situasi fiktif.','Caption')

for i,x in enumerate(items,1):
    doc.add_page_break();para(f'{i:02}  /  {x["tag"]}','Subtitle');doc.add_heading(x['title'],1)
    para(x['summary'],'Subtitle')
    doc.add_heading('Penjelasan',2);para(x['explanation'])
    label('Perawat:',x['nurse']);label('Pasien:',x['patient'])
    p=para('');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(7)
    if 'image' in x:
        path=ROOT/'speech-assets'/x['image']
        if i==6:p.add_run().add_picture(str(path),height=Cm(5.1))
        elif i==9:p.add_run().add_picture(str(path),height=Cm(5.0))
        else:p.add_run().add_picture(str(path),width=Cm(16.5))
        caption=x['caption']
    else:
        path=illustration(i);p.add_run().add_picture(str(path),width=Cm(16.5))
        refs=''.join(f'[{k}]' for k in x['refs'])
        caption=f'Ilustrasi penjelas yang dibuat khusus untuk dokumen ini; contoh fiktif. Dasar konsep: {refs}.'
    para(f'Gambar {i}. {caption}','Caption')
    doc.add_heading('Contoh interaksi',2);para(x['example'])
    label('Hal yang perlu dipahami:',x['note'])
    p=para('Sumber penjelasan: ')
    for idx,k in enumerate(x['refs']):
        if idx:p.add_run('  •  ')
        link(p,f'[{k}] '+sources[k-1][0].split('.')[0],sources[k-1][1])
    for r in p.runs:r.font.size=Pt(9)

for start,end in [(0,6),(6,11)]:
    doc.add_page_break();doc.add_heading('Daftar sumber'+(' — lanjutan' if start else ''),1)
    if start==0:
        para('Sumber diperiksa pada 16 Oktober 2026. Nomor rujukan pada setiap uraian dan keterangan gambar merujuk ke daftar ini. Judul dan tautan dapat diklik.')
    for j in range(start,end):
        title,url,desc=sources[j]
        p=para('');p.paragraph_format.space_before=Pt(10);p.paragraph_format.keep_with_next=True
        p.add_run(f'[{j+1}] {title}').bold=True
        p=para(desc);p.paragraph_format.keep_with_next=True
        p=para('');link(p,url,url)
        for r in p.runs:r.font.size=Pt(9)
    if start:
        doc.add_heading('Keterangan gambar dan rujukan',2)
        para('Gambar 1, 2, 3, 5, 7, dan 8 adalah ilustrasi penjelas asli. Gambar 4, 6, dan 9 merupakan cuplikan sumber yang disebutkan pada caption. Hak gambar rujukan tetap pada pemiliknya.')
        para('Artikel profesi, panduan layanan, dan dokumentasi produk memiliki tujuan berbeda. Rujukan di sini menjelaskan bentuk komunikasi dan contoh alat; tidak digunakan untuk menyimpulkan bahwa seluruh cara sama efektifnya bagi semua pasien.')

out=ROOT/'word/referensi/Dokumentasi-Speech-Perawat-ke-Pasien-Revisi-Diskusi.docx'
doc.save(out)
(ASSETS/'README.md').write_text('# Gambar bahan diskusi\n\nIlustrasi 01, 02, 03, 05, 07, dan 08 dibuat khusus untuk penjelasan umum, bukan tampilan aplikasi. Dasar konsep dan tautan sumber ada pada dokumen revisi.\n\nGambar sumber resmi menggunakan berkas rujukan di folder induk speech-assets; caption dokumen memuat atribusi.\n',encoding='utf-8')
(Path(__file__).parent/'content.json').write_text(json.dumps({'items':items,'sources':sources},ensure_ascii=False,indent=2),encoding='utf-8')
print(out)
print('Images:',len(doc.inline_shapes),'Source entries:',len(sources))
